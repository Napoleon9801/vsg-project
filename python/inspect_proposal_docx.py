"""Recon script: dump paragraph indices + tables of 66AOS02-แก้.docx."""
from docx import Document
import sys

DOCX = r"D:\workspace\engineering\vsg-project\docs\thesis\66AOS02-แก้.docx"

doc = Document(DOCX)

sys.stdout.reconfigure(encoding="utf-8")

print("=== PARAGRAPHS (idx, style, text[:140]) ===")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.replace("\n", " ").replace("\r", " ")
    style = p.style.name if p.style else "None"
    print(f"[{i}] ({style}) {txt[:140]}")

print()
print("=== TABLES ===")
for ti, t in enumerate(doc.tables):
    print(f"-- Table {ti} ({len(t.rows)} rows x {len(t.columns)} cols) --")
    for ri, row in enumerate(t.rows):
        cells = [c.text.replace("\n", " | ")[:80] for c in row.cells]
        print(f"  R{ri}: {cells}")
