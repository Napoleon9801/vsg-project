"""Export current thesis Markdown to separate dated review DOCX files.
Run with the Codex bundled Python runtime (python-docx) and installed Pandoc.
Original advisor DOCX files are never overwritten. Native Word rendering is a
separate QA step, not inferred from successful conversion.
"""
from pathlib import Path
import argparse
import json
import subprocess
import hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
CHAPTERS = ('chapter3_methodology', 'chapter4_simulation_results', 'chapter5_hardware_design')


def set_font(style, size):
    style.font.name = 'Tahoma'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    props = style.element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); props.append(fonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:'+key), 'Tahoma')
    cs = OxmlElement('w:szCs'); cs.set(qn('w:val'), str(int(size*2))); props.append(cs)


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--date', default='2026-10-13')
    parser.add_argument('--pandoc', default=r'C:\Program Files\Pandoc\pandoc.exe')
    args=parser.parse_args()
    output=ROOT/'docs/thesis'/('review_'+args.date)
    output.mkdir(parents=True,exist_ok=True)
    manifest=[]
    for name in CHAPTERS:
        source=ROOT/'docs/thesis'/(name+'.md')
        target=output/(name+'.docx')
        subprocess.run([args.pandoc,str(source),'-f','markdown+tex_math_dollars',
                        '-t','docx','--standalone','-o',str(target)],check=True,cwd=source.parent)
        doc=Document(target)
        for section in doc.sections:
            section.page_width=Inches(8.5); section.page_height=Inches(11)
            section.top_margin=Inches(.8); section.bottom_margin=Inches(.8)
            section.left_margin=Inches(.85); section.right_margin=Inches(.85)
            section.header_distance=Inches(.3); section.footer_distance=Inches(.3)
        for style in doc.styles:
            if style.type in (1,2):
                set_font(style,11)
        for name_style, size in [('Title',20),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
            set_font(doc.styles[name_style],size)
            doc.styles[name_style].paragraph_format.keep_with_next=True
        normal=doc.styles['Normal'].paragraph_format
        normal.line_spacing=1.12; normal.space_after=Pt(6)
        for para in doc.paragraphs:
            para.paragraph_format.widow_control=True
            following=para._p.getnext()
            if following is not None and following.tag==qn('w:tbl'):
                para.paragraph_format.keep_with_next=True
            if para.text.strip().startswith('(2)') and len(para.text)<180:
                para.paragraph_format.keep_with_next=True
        if doc.paragraphs:
            doc.paragraphs[0].style=doc.styles['Title']
        for table in doc.tables:
            table.autofit=True
            props=table._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                border=OxmlElement('w:'+edge)
                for key,value in (('val','single'),('sz','4'),('color','D9D9D9')):
                    border.set(qn('w:'+key),value)
                borders.append(border)
            props.append(borders)
            for row_index,row in enumerate(table.rows):
                trpr=row._tr.get_or_add_trPr()
                no_split=OxmlElement('w:cantSplit'); trpr.append(no_split)
                if row_index==0:
                    repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
                for cell in row.cells:
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell_props=cell._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    for edge in ('top','left','bottom','right'):
                        margin=OxmlElement('w:'+edge)
                        margin.set(qn('w:w'),'40' if edge in ('top','bottom') else '70')
                        margin.set(qn('w:type'),'dxa')
                        margins.append(margin)
                    cell_props.append(margins)
                    if row_index==0:
                        shading=OxmlElement('w:shd')
                        shading.set(qn('w:fill'),'E8EEF3')
                        cell_props.append(shading)
                    for para in cell.paragraphs:
                        para.paragraph_format.keep_with_next=(row_index==0)
                        para.paragraph_format.space_after=Pt(2)
                        para.paragraph_format.space_before=Pt(2)
                        for run in para.runs:
                            run.font.name='Tahoma'; run.font.size=Pt(10)
        footer=doc.sections[0].footer.paragraphs[0]
        footer.alignment=2
        footer.add_run('Review '+args.date+'  |  ')
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
        doc.core_properties.comments='Dated review export; not a replacement for advisor originals. Physical validation remains pending.'
        doc.save(target)
        reopened=Document(target)
        xml=reopened.element.xml
        manifest.append({'source':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
                         'output':str(target),'output_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
                         'paragraphs':len(reopened.paragraphs),'tables':len(reopened.tables),
                         'native_math_elements':xml.count('<m:oMath>'),'render_verified':False})
        print(target)
    (output/'export_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')


if __name__=='__main__':
    main()
