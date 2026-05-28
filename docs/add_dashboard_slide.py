"""
Add a Project Dashboard slide to 66AOS02_final.pptx.

The dashboard is a single information-dense slide that gives committee
members a 30-second overview of the entire project state:

  Row 1  : 4 KPI tiles  - Objectives | Verified Metrics | Design Point | ISR Budget
  Row 2  : Mini-Gantt timeline (16 weeks) with phase coloring
  Row 3  : Highlight strip with one-liner facts

Inserted at position 20 (between '6.2 Expected results' and 'End presentation').
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


SRC = Path(r"C:\Users\User\Downloads\66AOS02_final.pptx")
DST = Path(r"C:\Users\User\Downloads\66AOS02_final_v2.pptx")

FONT = "Tahoma"

# Palette - dashboard tile colors
COL_ORANGE   = RGBColor(0xE9, 0x6A, 0x1F)   # accent (SUT brand)
COL_GREEN    = RGBColor(0x2E, 0x7D, 0x32)   # status: done / safe
COL_GREEN_BG = RGBColor(0xE8, 0xF5, 0xE9)
COL_AMBER    = RGBColor(0xF5, 0x7C, 0x00)   # status: in-progress / watch
COL_AMBER_BG = RGBColor(0xFF, 0xF3, 0xE0)
COL_BLUE     = RGBColor(0x15, 0x65, 0xC0)
COL_BLUE_BG  = RGBColor(0xE3, 0xF2, 0xFD)
COL_CYAN_BG  = RGBColor(0xE0, 0xF7, 0xFA)
COL_CYAN     = RGBColor(0x00, 0x83, 0x8F)
COL_GREY_BG  = RGBColor(0xF5, 0xF5, 0xF5)
COL_GREY     = RGBColor(0x42, 0x42, 0x42)
COL_DARK     = RGBColor(0x21, 0x21, 0x21)
COL_WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
COL_RED      = RGBColor(0xC6, 0x28, 0x28)


# ----- low level helpers ----------------------------------------------------

def add_rounded_box(slide, x, y, w, h, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def add_rect(slide, x, y, w, h, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.5)
    shape.shadow.inherit = False
    return shape


def add_text(slide, x, y, w, h, text, *, size=14, bold=False,
             color=COL_DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font=FONT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def fit_text_in_shape(shape, text, *, size=14, bold=False,
                     color=COL_DARK, align=PP_ALIGN.LEFT,
                     anchor=MSO_ANCHOR.MIDDLE, font=FONT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = ""
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


# ----- KPI tile ------------------------------------------------------------

def add_kpi_tile(slide, x, y, w, h, *, header, header_color,
                 body_lines, big_value=None, big_unit=None,
                 tile_bg=COL_WHITE):
    """Add a KPI tile with a colored header bar, a big value (optional),
    and a list of body bullets."""
    # main tile background
    add_rounded_box(slide, x, y, w, h, tile_bg, line_color=COL_GREY)

    # header bar
    hdr_h = Inches(0.32)
    add_rect(slide, x, y, w, hdr_h, header_color)
    add_text(slide, x + Inches(0.10), y + Inches(0.03),
             w - Inches(0.20), hdr_h,
             header, size=11, bold=True, color=COL_WHITE,
             anchor=MSO_ANCHOR.MIDDLE)

    body_y = y + hdr_h + Inches(0.08)
    body_h = h - hdr_h - Inches(0.10)

    # big value block
    if big_value is not None:
        big_h = Inches(0.55)
        add_text(slide, x + Inches(0.10), body_y,
                 w - Inches(0.20), big_h,
                 big_value, size=24, bold=True,
                 color=header_color, anchor=MSO_ANCHOR.MIDDLE)
        if big_unit:
            add_text(slide, x + Inches(0.10),
                     body_y + Inches(0.55),
                     w - Inches(0.20), Inches(0.22),
                     big_unit, size=10, color=COL_GREY,
                     anchor=MSO_ANCHOR.TOP)
        body_y = body_y + Inches(0.85)

    # body bullets
    for line, line_color in body_lines:
        add_text(slide, x + Inches(0.10), body_y,
                 w - Inches(0.20), Inches(0.20),
                 line, size=10, color=line_color)
        body_y += Inches(0.21)


# ----- Gantt strip ---------------------------------------------------------

def add_mini_gantt(slide, x, y, w, h):
    """Render a single-bar mini-Gantt across 16 weeks with phase coloring."""
    # background
    add_rounded_box(slide, x, y, w, h, COL_WHITE, line_color=COL_GREY)

    # header
    add_text(slide, x + Inches(0.15), y + Inches(0.05),
             w - Inches(0.30), Inches(0.25),
             "แผนการดำเนินงาน 16 สัปดาห์  (Project Timeline)",
             size=11, bold=True, color=COL_DARK)

    # geometry of the timeline strip
    bar_x = x + Inches(0.15)
    bar_y = y + Inches(0.42)
    bar_w = w - Inches(0.30)
    bar_h = Inches(0.45)
    week_w = bar_w / 16

    # phases: (start_week, end_week, label, color, done_fraction)
    phases = [
        (1, 2,   "1. ทฤษฎี",       COL_GREEN, 1.0),   # done early
        (2, 6,   "2. Simulation",   COL_GREEN, 0.85),  # mostly done (Stages 1-5a)
        (5, 9,   "3. HW design",    COL_AMBER, 0.0),
        (9, 12,  "4. Firmware",     COL_AMBER, 0.0),
        (11, 14, "5. Integration",  COL_AMBER, 0.0),
        (14, 16, "6. Report",       COL_BLUE,  0.0),
    ]

    # base track (light grey)
    add_rect(slide, bar_x, bar_y, bar_w, bar_h, COL_GREY_BG,
             line_color=COL_GREY)

    # week tick labels every 4 weeks
    for wk in [1, 4, 8, 12, 16]:
        tick_x = bar_x + Emu(int(week_w * (wk - 1)))
        add_text(slide, tick_x - Inches(0.1),
                 bar_y + bar_h + Inches(0.03),
                 Inches(0.3), Inches(0.18),
                 f"W{wk}", size=8, color=COL_GREY,
                 align=PP_ALIGN.LEFT)

    # phase bars stacked, with phase label inside
    row_offset = [0, Emu(int(bar_h)) // 6, Emu(int(bar_h)) // 6 * 2,
                  Emu(int(bar_h)) // 6 * 3, Emu(int(bar_h)) // 6 * 4,
                  Emu(int(bar_h)) // 6 * 5]
    row_h = bar_h // 6
    for i, (start, end, label, color, done) in enumerate(phases):
        seg_x = bar_x + Emu(int(week_w * (start - 1)))
        seg_w = Emu(int(week_w * (end - start + 1)))
        ry = bar_y + row_offset[i]
        add_rect(slide, seg_x, ry, seg_w, row_h, color)
        if done > 0:
            done_w = Emu(int(int(seg_w) * done))
            add_rect(slide, seg_x, ry, done_w, row_h,
                     RGBColor(0x1B, 0x5E, 0x20))  # darker green = "done"
        add_text(slide, seg_x + Inches(0.04), ry,
                 seg_w, row_h, label, size=8, bold=True,
                 color=COL_WHITE, anchor=MSO_ANCHOR.MIDDLE)

    # "today" marker: pre-W1 (defense day)
    today_x = bar_x - Inches(0.08)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  today_x, bar_y - Inches(0.05),
                                  Emu(28000), bar_h + Inches(0.10))
    line.fill.solid(); line.fill.fore_color.rgb = COL_RED
    line.line.fill.background()
    add_text(slide, today_x - Inches(0.35),
             bar_y - Inches(0.25),
             Inches(0.85), Inches(0.20),
             "วันสอบ", size=8, bold=True, color=COL_RED)


# ----- highlight strip -----------------------------------------------------

def add_highlight_strip(slide, x, y, w, h):
    add_rounded_box(slide, x, y, w, h, COL_AMBER_BG,
                    line_color=COL_AMBER)
    add_text(slide, x + Inches(0.15), y + Inches(0.05),
             w - Inches(0.30), Inches(0.22),
             "Project Highlights  (Key Differentiators)",
             size=11, bold=True, color=COL_AMBER)
    items = [
        "💡  Innovation : ESP32-based low-cost GFM inverter @ 48 V / 100 W "
        "(literature uses TI C2000 / FPGA at 10x cost)",
        "✅  Sim Validated : Stages 1–5a all match analytical predictions "
        "to within 1.5 % error (V_d, P, Δf, ω_n, ζ)",
        "🛠️  Engineering Finding : Discovered overmodulation @ 24 V "
        "via simulation — revised design to 48 V before hardware build",
        "⚡  ESP32 Feasibility Proven : ISR @ 10 kHz uses only 8 % CPU — "
        "92 % headroom available for Adaptive VSG / fault detection",
    ]
    txt_y = y + Inches(0.32)
    for line in items:
        add_text(slide, x + Inches(0.18), txt_y,
                 w - Inches(0.30), Inches(0.24),
                 line, size=10, color=COL_DARK)
        txt_y += Inches(0.24)


# ----- assemble the slide --------------------------------------------------

def build_dashboard_slide(prs):
    layout = prs.slide_layouts[5]   # Title Only
    slide  = prs.slides.add_slide(layout)

    # Title bar
    title = slide.shapes.title
    title.text = ""
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "Project Dashboard  —  สถานะโครงงาน 66AOS02"
    run.font.name = FONT
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COL_ORANGE

    # Subtitle (below title)
    add_text(slide, Inches(0.55), Inches(0.95),
             Inches(9.0), Inches(0.25),
             "VSG for BESS Inverter Control  |  สอบหัวข้อ 28 พฤษภาคม 2569",
             size=11, color=COL_GREY)

    # ---- Row 1 : 4 KPI tiles ----
    row1_y = Inches(1.30)
    row1_h = Inches(1.85)
    margin = Inches(0.20)
    gap    = Inches(0.10)
    total_w = Inches(10.0) - 2 * margin - 3 * gap
    tile_w  = total_w / 4

    # Tile 1: Objectives
    add_kpi_tile(slide, margin, row1_y, tile_w, row1_h,
                 header="วัตถุประสงค์ (4 ข้อ)",
                 header_color=COL_BLUE,
                 tile_bg=COL_BLUE_BG,
                 big_value="1 / 4",
                 big_unit="ทำเสร็จเชิง simulation",
                 body_lines=[
                     ("✅ 1. VSG Simulink Model",                COL_GREEN),
                     ("⏳ 2. VSG vs Droop compare",              COL_GREY),
                     ("⏳ 3. H, D parameter sweep",              COL_GREY),
                     ("⏳ 4. Hardware prototype (ESP32)",        COL_GREY),
                 ])

    # Tile 2: Verified Metrics
    x2 = margin + tile_w + gap
    add_kpi_tile(slide, x2, row1_y, tile_w, row1_h,
                 header="Verified Sim Metrics  (Stage 5a)",
                 header_color=COL_GREEN,
                 tile_bg=COL_GREEN_BG,
                 big_value="< 1.5 %",
                 big_unit="error vs analytical",
                 body_lines=[
                     ("V_d  = 13.97 / 14.00 V  ✓",       COL_DARK),
                     ("P    = 137.9 / 139.6 W  ✓",        COL_DARK),
                     ("Δf   = −0.77 Hz  (settle 93%)",     COL_DARK),
                     ("ω_n, ζ, τ : match textbook",         COL_DARK),
                 ])

    # Tile 3: Design Point
    x3 = x2 + tile_w + gap
    add_kpi_tile(slide, x3, row1_y, tile_w, row1_h,
                 header="Design Point",
                 header_color=COL_ORANGE,
                 tile_bg=COL_AMBER_BG,
                 big_value="48 V / 100 W",
                 big_unit="ใน envelope ≤ 60 V / ≤ 300 W",
                 body_lines=[
                     ("Modulation index = 0.58 (linear)",   COL_DARK),
                     ("V_DC ห่าง MOSFET limit 80 V",         COL_DARK),
                     ("BESS : 4 × 12 V LiFePO4",            COL_DARK),
                     ("LC : 1.5 mH / 15 µF + R_d 1 Ω",      COL_DARK),
                 ])

    # Tile 4: ISR Budget
    x4 = x3 + tile_w + gap
    add_kpi_tile(slide, x4, row1_y, tile_w, row1_h,
                 header="ESP32 ISR Budget  @ 10 kHz",
                 header_color=COL_CYAN,
                 tile_bg=COL_CYAN_BG,
                 big_value="8 / 100 µs",
                 big_unit="ใช้ 8 %  /  headroom 92 %",
                 body_lines=[
                     ("Dual-core 240 MHz + FPU",            COL_DARK),
                     ("Core 0 : control ISR",                COL_DARK),
                     ("Core 1 : telemetry + WiFi",           COL_DARK),
                     ("รองรับ Adaptive H/D ในอนาคต",        COL_DARK),
                 ])

    # ---- Row 2 : Mini-Gantt ----
    row2_y = Inches(3.30)
    row2_h = Inches(1.30)
    add_mini_gantt(slide,
                   x=margin, y=row2_y,
                   w=Inches(10.0) - 2 * margin,
                   h=row2_h)

    # ---- Row 3 : Highlight strip ----
    row3_y = Inches(4.75)
    row3_h = Inches(2.50)
    add_highlight_strip(slide,
                        x=margin, y=row3_y,
                        w=Inches(10.0) - 2 * margin,
                        h=row3_h)

    return slide


# ----- slide reorder util ---------------------------------------------------

def move_slide_to(prs, src_idx, dst_idx):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    moved = slides[src_idx]
    xml_slides.remove(moved)
    xml_slides.insert(dst_idx, moved)


# ----- driver ---------------------------------------------------------------

def main():
    prs = Presentation(str(SRC))
    n0 = len(prs.slides)
    print(f"Loaded {SRC.name}: {n0} slides")

    print("Building Project Dashboard slide...")
    build_dashboard_slide(prs)

    # New slide is at idx n0 (end). Insert between '6.2 Expected results'
    # (idx 18) and 'จบการนำเสนอ' (idx 19). So target position = 19.
    move_slide_to(prs, src_idx=n0, dst_idx=19)

    prs.save(str(DST))
    print(f"Saved -> {DST}")
    print(f"Final deck: {len(prs.slides)} slides")
    print("\nFinal slide order:")
    for i, s in enumerate(prs.slides, 1):
        t = ""
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip():
                            t = r.text.strip(); break
                    if t: break
                if t: break
        print(f"  {i:2d}. {t[:78]}")


if __name__ == "__main__":
    main()
