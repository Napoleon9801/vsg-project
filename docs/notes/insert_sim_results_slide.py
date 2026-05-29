"""
Insert a new 'Preliminary Simulation Results' slide between the current
Test Plan slide (#14) and the Gantt slide (#15) in
proposal_defense_66AOS02.pptx, then save as a new file so the original
is preserved.

The new slide showcases that Stages 1-5a are already verified end-to-end:
this turns the deck from 'a plan' into 'a plan with proof' and directly
counters the innovation 3/5 rating in advisor feedback (2026-05-26 13:19).
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from lxml import etree


HERE   = Path(__file__).resolve().parent
SRC    = HERE / "proposal_defense_66AOS02.pptx"
DST    = HERE / "proposal_defense_66AOS02_v2.pptx"
PLOTS  = HERE.parent / "python" / "plots"
P4_IMG = PLOTS / "phase4_dq.png"
P5_IMG = PLOTS / "phase5a_switching.png"

INSERT_AFTER_INDEX = 13   # 0-based: insert after slide 14 (Test Plan)

THAI_FONT = "TH Sarabun New"
ORANGE    = RGBColor(0xE9, 0x6A, 0x1F)
DARKGREY  = RGBColor(0x33, 0x33, 0x33)
GREEN_OK  = RGBColor(0x2E, 0x7D, 0x32)


def add_paragraph(tf, text, *, size=14, bold=False, color=DARKGREY,
                  font=THAI_FONT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = ""
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return p


def build_slide(prs):
    """Build the new slide using the Title Only layout."""
    layout = prs.slide_layouts[5]   # Title Only
    slide  = prs.slides.add_slide(layout)

    # --- Title ---
    title_shape = slide.shapes.title
    title_shape.text = "ผลการจำลองเบื้องต้น (Preliminary Simulation Results)"
    for para in title_shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.name = THAI_FONT
            run.font.size = Pt(28)
            run.font.bold = True
            run.font.color.rgb = ORANGE

    # --- Left text block: key findings ---
    left_box = slide.shapes.add_textbox(
        left=Inches(0.35), top=Inches(1.25),
        width=Inches(4.6), height=Inches(5.9)
    )
    tf = left_box.text_frame
    tf.word_wrap = True

    add_paragraph(tf, "Stages 1-4 (Verified vs Analytical)",
                  size=18, bold=True, color=ORANGE, first=True)
    add_paragraph(tf, "• Stage 1 Swing Eq:  Δf = +0.0016 Hz  (match)", size=14)
    add_paragraph(tf, "• Stage 2 Power-angle:  ωn=7.5, ζ=0.62  (match)", size=14)
    add_paragraph(tf, "• Stage 3 Droop islanded:  Δf=−0.25 Hz, τ=0.5s  (match)", size=14)
    add_paragraph(tf, "• Stage 4 Inverter dq:  Vd → ref < 1 mV", size=14)

    add_paragraph(tf, " ", size=8)
    add_paragraph(tf, "Stage 5a — Full 3-phase Switching Layer",
                  size=18, bold=True, color=ORANGE)
    add_paragraph(tf, "• V_dc = 48 V, SPWM 10 kHz, m_index = 0.58", size=14)
    add_paragraph(tf, "• V_d  = 13.971 V  (ref 14.00, err < 0.03 V)", size=14, color=GREEN_OK)
    add_paragraph(tf, "• P    = 137.9 W   (pred 139.6, err < 1.5 %)", size=14, color=GREEN_OK)
    add_paragraph(tf, "• Δf   = −0.77 Hz @ t=1.5s  (settle ~93 %)", size=14, color=GREEN_OK)

    add_paragraph(tf, " ", size=8)
    add_paragraph(tf, "ข้อค้นพบเชิงวิศวกรรม (Engineering Finding)",
                  size=16, bold=True, color=ORANGE)
    add_paragraph(tf, "Stage 5a sim discovered overmodulation",
                  size=13)
    add_paragraph(tf, "at V_dc = 24 V (V_set = 14 V > V_dc/2 = 12 V)",
                  size=13)
    add_paragraph(tf, "→  design point revised to 48 V before",
                  size=13, bold=True)
    add_paragraph(tf, "    hardware build (within ≤60 V scope envelope)",
                  size=13, bold=True)

    # --- Right side: two stacked plots ---
    # phase4 plot (top)
    if P4_IMG.exists():
        slide.shapes.add_picture(
            str(P4_IMG),
            left=Inches(5.05), top=Inches(1.25),
            width=Inches(4.85)
        )
    # phase5a plot (bottom)
    if P5_IMG.exists():
        slide.shapes.add_picture(
            str(P5_IMG),
            left=Inches(5.05), top=Inches(4.10),
            width=Inches(4.85)
        )

    return slide


def reorder_slide(prs, src_index, dst_index):
    """Move the slide currently at src_index to position dst_index in the
    presentation's XML slide list."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    moved = slides[src_index]
    xml_slides.remove(moved)
    xml_slides.insert(dst_index, moved)


def main():
    prs = Presentation(str(SRC))
    n_before = len(prs.slides)
    print(f"Original deck: {n_before} slides")
    print(f"Inserting new slide after index {INSERT_AFTER_INDEX} "
          f"(current slide #{INSERT_AFTER_INDEX+1})")

    build_slide(prs)
    # the new slide is currently at the end (index n_before).
    # move it to position INSERT_AFTER_INDEX+1 so it sits between the
    # Test Plan slide and the Gantt slide.
    reorder_slide(prs, src_index=n_before, dst_index=INSERT_AFTER_INDEX + 1)

    prs.save(str(DST))
    print(f"Saved -> {DST}")
    print(f"Total slides now: {len(prs.slides)}")


if __name__ == "__main__":
    main()
