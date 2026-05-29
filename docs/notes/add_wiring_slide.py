"""
Insert wiring diagram slide into 66AOS02_final_v2.pptx.

Place: right after slide 13 (5.1 BoM) so the hardware story flows:
   13. 5.1 BoM (what components)
   14. NEW Wiring diagram (how they connect)
   15. 5.2 ISR Budget (how MCU handles them)
   ...

Saves to 66AOS02_final_v3.pptx so original v2 stays intact.
"""

from __future__ import annotations
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


SRC = Path(r"C:\Users\User\Downloads\66AOS02_final_v2.pptx")
DST = Path(r"C:\Users\User\Downloads\66AOS02_final_v3.pptx")
IMG = Path(r"C:\Users\User\workspace\engineering\vsg-project\python\plots\wiring_diagram.png")

FONT      = "Tahoma"
COL_OR    = RGBColor(0xE9, 0x6A, 0x1F)
COL_DARK  = RGBColor(0x21, 0x21, 0x21)
COL_GREY  = RGBColor(0x55, 0x55, 0x55)


def build_wiring_slide(prs):
    layout = prs.slide_layouts[5]   # Title Only
    slide  = prs.slides.add_slide(layout)

    # Title
    title = slide.shapes.title
    title.text = ""
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "5.1.1  แผนผังต่ออุปกรณ์ฮาร์ดแวร์  (Hardware Wiring Diagram)"
    run.font.name = FONT
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COL_OR

    # Embed the PNG (large, centered)
    if IMG.exists():
        # slide is 10x7.5 inch. Title ~0.7 in. Image ~6.5 in tall, full width.
        slide.shapes.add_picture(
            str(IMG),
            left=Inches(0.20), top=Inches(0.95),
            width=Inches(9.60)
        )

    # Footnote
    foot = slide.shapes.add_textbox(
        Inches(0.30), Inches(7.05),
        Inches(9.40), Inches(0.35)
    )
    p = foot.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = ("Power chain (แดง+ดำ)  |  Sensing (น้ำเงิน)  |  PWM/control (ม่วง)  |  "
                "Protection (เส้นแดงเข้ม)  —  detail ใน Chapter 5 Hardware Design")
    run.font.name = FONT
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = COL_GREY

    return slide


def move_slide_to(prs, src_idx, dst_idx):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    moved = slides[src_idx]
    xml_slides.remove(moved)
    xml_slides.insert(dst_idx, moved)


def main():
    prs = Presentation(str(SRC))
    n0 = len(prs.slides)
    print(f"Loaded {SRC.name}: {n0} slides")

    print("Building wiring diagram slide...")
    build_wiring_slide(prs)

    # New slide is at idx n0 (end). Move to idx 13 (after BoM at idx 12).
    move_slide_to(prs, src_idx=n0, dst_idx=13)

    prs.save(str(DST))
    print(f"Saved -> {DST}")
    print(f"Total slides: {len(prs.slides)}")

    print("\nFinal slide order:")
    for i, s in enumerate(prs.slides, 1):
        t = ""
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip():
                            t = r.text.strip(); break
                    if t: break
                if t: break
        print(f"  {i:2d}. {t[:78]}")


if __name__ == "__main__":
    main()
